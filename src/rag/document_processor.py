"""
Document processing and chunking utilities.
"""

import os
from typing import List, Dict, Any, Iterable, Optional
from pathlib import Path
import hashlib
from ..utils.logger import get_logger
from ..config import config

logger = get_logger(__name__)


class DocumentProcessor:
    """Process and chunk documents for RAG"""
    
    def __init__(self, chunk_size: int = None, chunk_overlap: int = None):
        self.chunk_size = chunk_size or config.chunk_size
        self.chunk_overlap = chunk_overlap or config.chunk_overlap
        self.max_chunks_per_document = max(1, config.max_chunks_per_document)
    
    def process_text_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Read and process a text file"""
        try:
            file_name = os.path.basename(file_path)
            file_hash = self._hash_file(file_path)
            chunks = list(self._chunk_text_stream(file_path))
            documents = self._build_documents(
                chunks=chunks,
                file_name=file_name,
                file_type="text",
                file_hash=file_hash
            )
            logger.info(f"Processed text file {file_name} into {len(documents)} chunks")
            return documents
        except Exception as e:
            logger.error(f"Error processing text file: {e}")
            return []
    
    def process_pdf_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Process a PDF file with text, table, and optional image OCR extraction."""
        try:
            from pypdf import PdfReader
            
            reader = PdfReader(file_path)
            file_name = os.path.basename(file_path)
            file_hash = self._hash_file(file_path)
            page_texts: List[str] = []

            table_data = self._extract_pdf_tables(file_path) if config.pdf_extract_tables else {}
            image_data = self._extract_pdf_image_text(file_path) if config.pdf_extract_images_text else {}

            for page_num, page in enumerate(reader.pages):
                try:
                    section_parts: List[str] = []
                    extracted_text = (page.extract_text() or "").strip()
                    if extracted_text:
                        section_parts.append(extracted_text)

                    if page_num in table_data and table_data[page_num]:
                        section_parts.append(table_data[page_num])

                    if page_num in image_data and image_data[page_num]:
                        section_parts.append(image_data[page_num])

                    if section_parts:
                        page_texts.append(
                            f"--- Page {page_num + 1} ---\n" + "\n\n".join(section_parts)
                        )
                except Exception as e:
                    logger.warning(f"Error extracting text from page {page_num}: {e}")
            
            content = "\n\n".join(page_texts).strip()
            if not content:
                logger.warning(f"No text extracted from PDF: {file_path}")
                return []
            
            chunks = self.chunk_text(content)
            documents = self._build_documents(
                chunks=chunks,
                file_name=file_name,
                file_type="pdf",
                file_hash=file_hash
            )
            logger.info(f"Processed PDF {file_name} into {len(documents)} chunks")
            return documents
        except ImportError:
            logger.error("pypdf not installed. Install with: pip install -r requirements.txt")
            return []
        except Exception as e:
            logger.error(f"Error processing PDF file: {e}")
            return []
    
    def process_markdown_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Process a markdown file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            # Split by headers for better structure preservation
            chunks = self.chunk_markdown(content)
            file_name = os.path.basename(file_path)
            file_hash = self._hash_file(file_path)
            documents = self._build_documents(
                chunks=chunks,
                file_name=file_name,
                file_type="markdown",
                file_hash=file_hash
            )
            
            logger.info(f"Processed {file_name} into {len(documents)} chunks")
            return documents
        except Exception as e:
            logger.error(f"Error processing markdown file: {e}")
            return []
    
    def process_csv_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Process a CSV file"""
        try:
            import csv
            
            file_name = os.path.basename(file_path)
            rows = []
            
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                if reader.fieldnames is None:
                    logger.warning(f"CSV file {file_name} is empty or has no headers")
                    return []
                
                for row_num, row in enumerate(reader, start=2):  # start=2 to account for header row
                    # Convert row to readable text format
                    row_text = " | ".join([f"{k}: {v}" for k, v in row.items() if v])
                    rows.append(row_text)
            
            if not rows:
                logger.warning(f"No data found in CSV file: {file_name}")
                return []
            
            # Combine rows with header info
            header_row = " | ".join(reader.fieldnames) if reader.fieldnames else "CSV Data"
            content = f"CSV Headers: {header_row}\n\n" + "\n".join(rows)
            
            chunks = self.chunk_text(content)
            file_hash = self._hash_file(file_path)
            documents = self._build_documents(
                chunks=chunks,
                file_name=file_name,
                file_type="csv",
                file_hash=file_hash
            )
            
            logger.info(f"Processed CSV {file_name} into {len(documents)} chunks")
            return documents
        except Exception as e:
            logger.error(f"Error processing CSV file: {e}")
            return []
    
    def process_xlsx_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Process an XLSX (Excel) file"""
        try:
            from openpyxl import load_workbook
            
            file_name = os.path.basename(file_path)
            workbook = load_workbook(file_path)
            all_content = []
            
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                sheet_text = f"Sheet: {sheet_name}\n"
                
                # Get headers
                headers = []
                for cell in worksheet[1]:
                    headers.append(str(cell.value) if cell.value else "")
                
                sheet_text += "Headers: " + " | ".join(headers) + "\n\n"
                
                # Get data rows
                for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
                    row_text = " | ".join([f"{headers[i]}: {str(val)}" if i < len(headers) else str(val) 
                                          for i, val in enumerate(row) if val is not None])
                    if row_text.strip():
                        sheet_text += row_text + "\n"
                
                all_content.append(sheet_text)
            
            if not all_content or not any(c.strip() for c in all_content):
                logger.warning(f"No data found in XLSX file: {file_name}")
                return []
            
            content = "\n".join(all_content)
            chunks = self.chunk_text(content)
            file_hash = self._hash_file(file_path)
            documents = self._build_documents(
                chunks=chunks,
                file_name=file_name,
                file_type="xlsx",
                file_hash=file_hash
            )
            
            logger.info(f"Processed XLSX {file_name} into {len(documents)} chunks")
            return documents
        except ImportError:
            logger.error("openpyxl not installed. Install with: pip install -r requirements.txt")
            return []
        except Exception as e:
            logger.error(f"Error processing XLSX file: {e}")
            return []
    
    def process_xls_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Process an XLS (Excel) file"""
        try:
            import xlrd
            
            file_name = os.path.basename(file_path)
            workbook = xlrd.open_workbook(file_path)
            all_content = []
            
            for sheet_idx, worksheet in enumerate(workbook.sheets()):
                sheet_text = f"Sheet: {worksheet.name}\n"
                
                # Get headers
                headers = []
                if worksheet.nrows > 0:
                    for col_idx in range(worksheet.ncols):
                        headers.append(str(worksheet.cell_value(0, col_idx)))
                
                sheet_text += "Headers: " + " | ".join(headers) + "\n\n"
                
                # Get data rows
                for row_idx in range(1, worksheet.nrows):
                    row_values = []
                    for col_idx in range(worksheet.ncols):
                        cell_value = worksheet.cell_value(row_idx, col_idx)
                        if cell_value:
                            row_values.append(f"{headers[col_idx]}: {str(cell_value)}")
                    
                    if row_values:
                        sheet_text += " | ".join(row_values) + "\n"
                
                all_content.append(sheet_text)
            
            if not all_content or not any(c.strip() for c in all_content):
                logger.warning(f"No data found in XLS file: {file_name}")
                return []
            
            content = "\n".join(all_content)
            chunks = self.chunk_text(content)
            file_hash = self._hash_file(file_path)
            documents = self._build_documents(
                chunks=chunks,
                file_name=file_name,
                file_type="xls",
                file_hash=file_hash
            )
            
            logger.info(f"Processed XLS {file_name} into {len(documents)} chunks")
            return documents
        except ImportError:
            logger.error("xlrd not installed. Install with: pip install -r requirements.txt")
            return []
        except Exception as e:
            logger.error(f"Error processing XLS file: {e}")
            return []
    
    def process_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Process any supported file type
        
        Supported formats:
        - .txt: Plain text files
        - .pdf: PDF documents
        - .md, .markdown: Markdown files
        - .csv: Comma-separated values
        - .xls: Excel 97-2003 workbook
        - .xlsx: Excel 2007+ workbook
        """
        file_ext = Path(file_path).suffix.lower()
        
        if file_ext == '.pdf':
            return self.process_pdf_file(file_path)
        elif file_ext in ['.md', '.markdown']:
            return self.process_markdown_file(file_path)
        elif file_ext == '.csv':
            return self.process_csv_file(file_path)
        elif file_ext == '.xlsx':
            return self.process_xlsx_file(file_path)
        elif file_ext == '.xls':
            return self.process_xls_file(file_path)
        else:
            # Default to text file processing
            return self.process_text_file(file_path)
    
    def chunk_text(self, text: str) -> List[str]:
        """Split text into overlapping chunks"""
        if not text:
            return []
        chunks = []
        step = self.chunk_size - self.chunk_overlap
        if step <= 0:
            raise ValueError("chunk_size must be greater than chunk_overlap")
        
        for i in range(0, len(text), step):
            chunk = text[i:i + self.chunk_size]
            if chunk.strip():
                chunks.append(chunk)
            if len(chunks) >= self.max_chunks_per_document:
                logger.warning(
                    f"Reached max chunks per document ({self.max_chunks_per_document}); truncating extra chunks."
                )
                break
        
        return chunks
    
    def chunk_markdown(self, text: str) -> List[str]:
        """Split markdown into chunks by headers when possible"""
        if not text:
            return []
        chunks = []
        current_chunk = ""
        
        lines = text.split('\n')
        for line in lines:
            # Check if line is a header
            if line.startswith('#'):
                if current_chunk and len(current_chunk) > self.chunk_size // 2:
                    chunks.append(current_chunk.strip())
                    if len(chunks) >= self.max_chunks_per_document:
                        logger.warning(
                            f"Reached max chunks per document ({self.max_chunks_per_document}) while chunking markdown."
                        )
                        return chunks
                    current_chunk = line + "\n"
                else:
                    current_chunk += line + "\n"
            else:
                current_chunk += line + "\n"
                
                # Split if chunk is too large
                if len(current_chunk) > self.chunk_size:
                    chunks.append(current_chunk.strip())
                    if len(chunks) >= self.max_chunks_per_document:
                        logger.warning(
                            f"Reached max chunks per document ({self.max_chunks_per_document}) while chunking markdown."
                        )
                        return chunks
                    current_chunk = ""
        
        if current_chunk.strip():
            chunks.append(current_chunk.strip())
            if len(chunks) > self.max_chunks_per_document:
                chunks = chunks[:self.max_chunks_per_document]
        
        return chunks
    
    def _build_documents(
        self,
        chunks: List[str],
        file_name: str,
        file_type: str,
        file_hash: str,
        extra_metadata: Optional[Dict[str, Any]] = None
    ) -> List[Dict[str, Any]]:
        documents = []
        total_chunks = len(chunks)
        if total_chunks == 0:
            return documents
        
        for i, chunk in enumerate(chunks):
            if not chunk.strip():
                continue
            doc_id = self._generate_doc_id(file_name, i, chunk, file_hash=file_hash)
            metadata = {
                "source": file_name,
                "chunk_index": i,
                "total_chunks": total_chunks,
                "file_type": file_type,
                "file_hash": file_hash
            }
            if extra_metadata:
                metadata.update(extra_metadata)
            documents.append(
                {
                    "id": doc_id,
                    "content": chunk,
                    "metadata": metadata
                }
            )
        
        return documents
    
    def _chunk_text_stream(self, file_path: str) -> Iterable[str]:
        """Stream a large text file into chunks with overlap."""
        step = self.chunk_size - self.chunk_overlap
        if step <= 0:
            raise ValueError("chunk_size must be greater than chunk_overlap")
        
        buffer = ""
        emitted = 0
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            while True:
                data = f.read(self.chunk_size * 4)
                if not data:
                    break
                buffer += data
                while len(buffer) >= self.chunk_size:
                    chunk = buffer[:self.chunk_size]
                    if chunk.strip():
                        yield chunk
                        emitted += 1
                    if emitted >= self.max_chunks_per_document:
                        logger.warning(
                            f"Reached max chunks per document ({self.max_chunks_per_document}); stopping stream chunking."
                        )
                        return
                    buffer = buffer[step:]
        
        if buffer.strip() and emitted < self.max_chunks_per_document:
            yield buffer
    
    def _extract_pdf_tables(self, file_path: str) -> Dict[int, str]:
        """Extract tables from PDF pages using pdfplumber if installed."""
        tables_by_page: Dict[int, str] = {}
        try:
            import pdfplumber
        except ImportError:
            logger.warning("pdfplumber is not installed; skipping PDF table extraction.")
            return tables_by_page
        
        try:
            with pdfplumber.open(file_path) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    page_tables = page.extract_tables() or []
                    extracted_rows = []
                    for table in page_tables:
                        for row in table:
                            if row:
                                cells = [str(cell).strip() if cell is not None else "" for cell in row]
                                extracted_rows.append(" | ".join(cells))
                    if extracted_rows:
                        tables_by_page[page_idx] = "Tables:\n" + "\n".join(extracted_rows)
        except Exception as e:
            logger.warning(f"PDF table extraction failed for {file_path}: {e}")
        
        return tables_by_page
    
    def _extract_pdf_image_text(self, file_path: str) -> Dict[int, str]:
        """Extract image OCR text from PDF pages if OCR dependencies are available and enabled."""
        image_text_by_page: Dict[int, str] = {}
        if not config.pdf_ocr_enabled:
            return image_text_by_page
        
        try:
            import fitz  # PyMuPDF
            from PIL import Image
            import pytesseract
        except ImportError:
            logger.warning("OCR dependencies missing (PyMuPDF/Pillow/pytesseract); skipping OCR extraction.")
            return image_text_by_page
        
        try:
            doc = fitz.open(file_path)
            for page_idx in range(len(doc)):
                page = doc[page_idx]
                ocr_text_segments: List[str] = []
                images = page.get_images(full=True)
                for img_ref in images:
                    xref = img_ref[0]
                    base_image = doc.extract_image(xref)
                    image_bytes = base_image.get("image")
                    if not image_bytes:
                        continue
                    try:
                        from io import BytesIO
                        pil_img = Image.open(BytesIO(image_bytes))
                        ocr_text = (pytesseract.image_to_string(pil_img) or "").strip()
                        if ocr_text:
                            ocr_text_segments.append(ocr_text)
                    except Exception as e:
                        logger.debug(f"OCR failed on an image in page {page_idx + 1}: {e}")
                if ocr_text_segments:
                    image_text_by_page[page_idx] = "Image OCR:\n" + "\n\n".join(ocr_text_segments)
            doc.close()
        except Exception as e:
            logger.warning(f"PDF image OCR extraction failed for {file_path}: {e}")
        
        return image_text_by_page
    
    @staticmethod
    def _hash_file(file_path: str) -> str:
        """Compute a stable hash for file-level dedupe metadata."""
        hasher = hashlib.sha256()
        with open(file_path, "rb") as f:
            for block in iter(lambda: f.read(1024 * 1024), b""):
                hasher.update(block)
        return hasher.hexdigest()
    
    @staticmethod
    def _generate_doc_id(filename: str, chunk_idx: int, content: str, file_hash: str = "") -> str:
        """Generate a unique document ID"""
        content_hash = hashlib.md5(content.encode()).hexdigest()[:8]
        source_hash = file_hash[:12] if file_hash else "nofilehash"
        return f"{filename}_{source_hash}_{chunk_idx}_{content_hash}"
